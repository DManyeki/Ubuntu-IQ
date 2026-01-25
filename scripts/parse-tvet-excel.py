# Parse TVET Excel file directly
# Usage: python scripts/parse-tvet-excel.py

import openpyxl
import json
import re

EXCEL_PATH = "Kuccps Files/TVET_CLUSTER_DOCUMENT_2025.xlsx"
OUTPUT_PATH = "scripts/data/tvet_excel_parsed.json"

def clean_grade(grade_str):
    """Normalize grade format"""
    if not grade_str:
        return ""
    grade = str(grade_str).strip()
    grade = re.sub(r'\s*\(plain\)', '', grade, flags=re.IGNORECASE)
    grade = re.sub(r'\s*\(minus\)', '-', grade, flags=re.IGNORECASE)
    grade = re.sub(r'\s*\(plus\)', '+', grade, flags=re.IGNORECASE)
    grade = grade.replace('C plain', 'C').replace('D plain', 'D')
    return grade.strip()

def parse_requirements(req_text):
    """Parse requirement text into structured list"""
    if not req_text or str(req_text).strip().lower() == 'none':
        return []
    
    req_str = str(req_text).strip()
    requirements = []
    
    # Split by common patterns
    parts = re.split(r'(?<=[A-Z+\-\)])\s+(?=[A-Z])', req_str)
    
    for part in parts:
        part = part.strip()
        if not part:
            continue
        
        # Skip section headers
        if any(x in part.lower() for x in ['science based', 'non-science', 'teaching subjects']):
            continue
        
        # Match "SUBJECT - GRADE" pattern
        match = re.match(r'^([A-Z][A-Z0-9/\s&]*(?:Alternative\s*[A-Z])?)\s*[-–]\s*(.+)$', part, re.IGNORECASE)
        if match:
            subject = match.group(1).strip()
            grade = clean_grade(match.group(2).split()[0] if ' ' in match.group(2) else match.group(2))
            if subject and grade and len(grade) <= 3:
                requirements.append({"subject": subject, "grade": grade})
    
    return requirements

def detect_exam_type(cell_value, current_type):
    """Detect exam type from cell content"""
    if cell_value:
        val = str(cell_value).upper()
        if 'KNEC' in val:
            return 'KNEC'
        if 'INTERNAL' in val:
            return 'INTERNAL'
    return current_type

def parse_excel():
    print("TVET EXCEL PARSER")
    print("=" * 50)
    
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    print(f"Sheets: {wb.sheetnames}")
    
    programs = []
    current_exam_type = 'KUCCPS'
    current_category = None
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\nProcessing {sheet_name}...")
        
        for row_idx in range(1, ws.max_row + 1):
            # Get cell values
            col_a = ws.cell(row=row_idx, column=1).value  # Number or exam type
            col_b = ws.cell(row=row_idx, column=2).value  # Category
            col_c = ws.cell(row=row_idx, column=3).value  # Level
            col_d = ws.cell(row=row_idx, column=4).value  # Mean Grade
            col_e = ws.cell(row=row_idx, column=5).value  # Requirements
            
            # Check for exam type headers
            current_exam_type = detect_exam_type(col_a, current_exam_type)
            current_exam_type = detect_exam_type(col_b, current_exam_type)
            
            # Skip header rows
            if col_b and 'Programme' in str(col_b):
                continue
            if col_e and 'Minimum Subject' in str(col_e):
                continue
            
            # Update category if provided
            if col_b and str(col_b).strip():
                cat = str(col_b).strip()
                if not any(x in cat for x in ['KNEC', 'INTERNAL', 'Programme']):
                    current_category = cat
            
            # Get level
            level = str(col_c).strip() if col_c else None
            if not level or 'Level' in level:
                continue
            
            # Clean level
            if 'Diploma' in level:
                level = 'Diploma'
            elif 'Certificate' in level:
                level = 'Certificate'
            elif 'Artisan' in level:
                level = 'Artisan'
            else:
                continue
            
            # Get mean grade
            mean_grade = clean_grade(col_d) if col_d else ''
            
            # Parse requirements
            requirements = parse_requirements(col_e)
            
            if current_category and level:
                program = {
                    "category": current_category,
                    "level": level,
                    "min_mean_grade": mean_grade,
                    "requirements": requirements,
                    "exam_type": current_exam_type
                }
                programs.append(program)
                print(f"  Found: {current_category} - {level}")
    
    # Deduplicate
    unique = {}
    for p in programs:
        key = f"{p['category']}|{p['level']}|{p['exam_type']}"
        if key not in unique or len(p['requirements']) > len(unique[key]['requirements']):
            unique[key] = p
    
    final_programs = list(unique.values())
    
    print(f"\n{'=' * 50}")
    print(f"Extracted {len(final_programs)} unique programs")
    
    # Save
    output = {
        "source": EXCEL_PATH,
        "extraction_method": "excel_parser",
        "total_programs": len(final_programs),
        "programs": final_programs
    }
    
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(output, f, indent=2, ensure_ascii=False)
    
    print(f"Saved to: {OUTPUT_PATH}")
    
    # Summary
    print("\nBy exam type:")
    for exam_type in ['KUCCPS', 'KNEC', 'INTERNAL']:
        count = len([p for p in final_programs if p['exam_type'] == exam_type])
        print(f"  {exam_type}: {count}")
    
    print("\nSample programs:")
    for p in final_programs[:5]:
        print(f"  {p['category']} ({p['level']}): {p['min_mean_grade']} [{p['exam_type']}]")

if __name__ == "__main__":
    parse_excel()
