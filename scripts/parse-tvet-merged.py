# Parse cleaned merged TVET Excel file
# Usage: python scripts/parse-tvet-merged.py

import openpyxl
import json
import re

EXCEL_PATH = "Kuccps Files/TVET_CLUSTER_DOCUMENT_2025 - updated.xlsx"
OUTPUT_PATH = "scripts/data/tvet_merged_parsed.json"

def clean_grade(grade_str):
    """Normalize grade format"""
    if not grade_str:
        return ""
    grade = str(grade_str).strip()
    # Handle various formats
    grade = re.sub(r'\s*\(plain\)', '', grade, flags=re.IGNORECASE)
    grade = re.sub(r'\s*plain', '', grade, flags=re.IGNORECASE)
    grade = re.sub(r'\s*\(minus\)', '-', grade, flags=re.IGNORECASE)
    grade = re.sub(r'\s*\(plus\)', '+', grade, flags=re.IGNORECASE)
    grade = grade.replace('C-minus', 'C-').replace('D-minus', 'D-')
    grade = grade.strip()
    # Extract just the grade letter
    match = re.match(r'^([A-D][+-]?)', grade)
    return match.group(1) if match else grade

def parse_requirements(req_text):
    """Parse requirements into structured list"""
    if not req_text:
        return []
    
    req_str = str(req_text).strip()
    if req_str.lower() in ['none', 'n/a', '-', '']:
        return []
    
    requirements = []
    
    # Split by common separators
    # Handle newlines, commas, and multiple spaces
    lines = re.split(r'[\n,]|\s{2,}', req_str)
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        # Skip headers/notes
        if any(x in line.lower() for x in ['science based', 'non-science', 'teaching', 'alternative a', 'alternative b']):
            if 'MATH Alternative' not in line and 'MAT Alternative' not in line:
                continue
        
        # Match patterns like "ENG/KIS - C+" or "MATH - C (plain)" or "BIO/CHE - C"
        match = re.match(r'^([A-Z][A-Z0-9/\s]*(?:\s*Alternative\s*[A-B])?)\s*[-–:]\s*(.+)$', line, re.IGNORECASE)
        if match:
            subject = match.group(1).strip()
            grade = clean_grade(match.group(2))
            
            # Validate
            if subject and grade and len(grade) <= 3 and grade[0] in 'ABCD':
                requirements.append({"subject": subject, "grade": grade})
    
    return requirements

def detect_exam_type(text):
    """Detect exam type from text"""
    if text:
        t = str(text).upper()
        if 'KNEC' in t:
            return 'KNEC'
        if 'INTERNAL' in t:
            return 'INTERNAL'
    return None

def parse_merged_excel():
    print("TVET MERGED EXCEL PARSER")
    print("=" * 50)
    
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active
    
    print(f"Sheet: {ws.title}")
    print(f"Rows: {ws.max_row}, Columns: {ws.max_column}")
    
    programs = []
    current_exam_type = 'KUCCPS'
    current_category = None
    
    for row_idx in range(1, ws.max_row + 1):
        # Get cell values (A=1, B=2, C=3, D=4, E=5)
        col_a = ws.cell(row=row_idx, column=1).value
        col_b = ws.cell(row=row_idx, column=2).value  # Category
        col_c = ws.cell(row=row_idx, column=3).value  # Level
        col_d = ws.cell(row=row_idx, column=4).value  # Mean Grade
        col_e = ws.cell(row=row_idx, column=5).value  # Requirements
        
        # Check for exam type markers
        for cell in [col_a, col_b]:
            exam = detect_exam_type(cell)
            if exam:
                current_exam_type = exam
                print(f"  → Switched to {exam}")
        
        # Skip header rows
        if col_b and 'Programme' in str(col_b):
            continue
        if col_e and 'Subject Grade' in str(col_e):
            continue
        
        # Update category if provided
        if col_b and str(col_b).strip():
            cat = str(col_b).strip()
            # Skip if it's an exam type header
            if not detect_exam_type(cat) and not 'Programme' in cat:
                current_category = cat
        
        # Get level
        level_raw = str(col_c).strip() if col_c else ""
        level = None
        if 'Diploma' in level_raw:
            level = 'Diploma'
        elif 'Certificate' in level_raw:
            level = 'Certificate'
        elif 'Artisan' in level_raw:
            level = 'Artisan'
        
        if not level:
            continue
        
        # Get mean grade
        mean_grade = clean_grade(col_d) if col_d else ""
        
        # Parse requirements
        requirements = parse_requirements(col_e)
        
        if current_category and level:
            program = {
                "category": current_category,
                "level": level,
                "min_mean_grade": mean_grade,
                "requirements": requirements,
                "exam_type": current_exam_type
            }
            programs.append(program)
            req_count = len(requirements)
            print(f"  Found: {current_category[:30]} - {level} ({mean_grade}) [{req_count} reqs]")
    
    # Deduplicate by key
    unique = {}
    for p in programs:
        key = f"{p['category']}|{p['level']}|{p['exam_type']}"
        if key not in unique or len(p['requirements']) > len(unique[key]['requirements']):
            unique[key] = p
    
    final = list(unique.values())
    
    print(f"\n{'=' * 50}")
    print(f"Total programs: {len(final)}")
    
    # Save
    output = {
        "source": EXCEL_PATH,
        "extraction_method": "merged_excel_parser",
        "total_programs": len(final),
        "programs": final
    }
    
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(output, f, indent=2, ensure_ascii=False)
    
    print(f"Saved to: {OUTPUT_PATH}")
    
    # Summary by type
    print("\nBy exam type:")
    for t in ['KUCCPS', 'KNEC', 'INTERNAL']:
        count = len([p for p in final if p['exam_type'] == t])
        print(f"  {t}: {count}")
    
    # Sample with requirements
    print("\nSample programs with requirements:")
    for p in [x for x in final if x['requirements']][:5]:
        print(f"  {p['category'][:35]} ({p['level']}): {p['min_mean_grade']}")
        for r in p['requirements'][:3]:
            print(f"    - {r['subject']}: {r['grade']}")

if __name__ == "__main__":
    parse_merged_excel()
