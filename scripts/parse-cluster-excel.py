# Parse cluster data from Excel file
# The Excel has multiple sheets (one per page), some inherit headers from previous page
# Usage: python scripts/parse-cluster-excel.py

import json
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl not installed. Run: pip install openpyxl")
    exit(1)

def parse_cluster_excel(excel_path):
    """Parse cluster data from Excel file with multiple sheets"""
    print(f"Loading Excel: {excel_path}")
    
    wb = openpyxl.load_workbook(excel_path)
    print(f"Found {len(wb.sheetnames)} sheets")
    
    clusters = []
    current_cluster_id = None
    current_subjects = {
        "subject_1": "", "grade_1": "",
        "subject_2": "", "grade_2": "",
        "subject_3": "", "grade_3": "",
        "subject_4": "", "grade_4": ""
    }
    current_programs = []
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\n  Processing: {sheet_name}")
        
        rows = list(sheet.iter_rows(values_only=True))
        
        for row_idx, row in enumerate(rows):
            if not row or all(c is None for c in row):
                continue
            
            # Get first cell value
            first_cell = str(row[0] or "").strip()
            
            # Check if this is a cluster header (e.g., "1", "2A", "5B", "13A")
            import re
            cluster_match = re.match(r'^(\d+[A-Z]?)$', first_cell)
            
            if cluster_match:
                # Save previous cluster if exists
                if current_cluster_id and current_programs:
                    clusters.append({
                        "cluster_id": current_cluster_id,
                        "subjects": current_subjects.copy(),
                        "programs": current_programs.copy()
                    })
                
                current_cluster_id = cluster_match.group(1)
                current_programs = []
                
                # Parse subjects from remaining cells
                cells = [str(c or "").strip() for c in row[1:] if c]
                
                # Reset subjects
                current_subjects = {
                    "subject_1": "", "grade_1": "",
                    "subject_2": "", "grade_2": "",
                    "subject_3": "", "grade_3": "",
                    "subject_4": "", "grade_4": ""
                }
                
                # Parse subject requirements from cells
                for i, cell in enumerate(cells[:4]):
                    subj_key = f"subject_{i+1}"
                    grade_key = f"grade_{i+1}"
                    
                    # Parse format like "MAT ALTERNATIVE A - C+" or "PHY - C (PLAIN)"
                    if " - " in cell:
                        parts = cell.rsplit(" - ", 1)
                        current_subjects[subj_key] = parts[0].strip()
                        current_subjects[grade_key] = parts[1].strip()
                    elif "–" in cell:  # Different dash character
                        parts = cell.rsplit("–", 1)
                        current_subjects[subj_key] = parts[0].strip()
                        current_subjects[grade_key] = parts[1].strip()
                    else:
                        current_subjects[subj_key] = cell
                        current_subjects[grade_key] = ""
                
                print(f"    Cluster {current_cluster_id}: {cells[0][:30] if cells else ''}...")
                continue
            
            # Check if this row contains program names
            for cell in row:
                if cell:
                    cell_text = str(cell).strip()
                    if cell_text.startswith(("Bachelor", "B.", "BACHELOR")):
                        # Clean and add program
                        current_programs.append(cell_text)
    
    # Don't forget last cluster
    if current_cluster_id and current_programs:
        clusters.append({
            "cluster_id": current_cluster_id,
            "subjects": current_subjects.copy(),
            "programs": current_programs.copy()
        })
    
    print(f"\nExtracted {len(clusters)} clusters")
    return clusters

def main():
    print("CLUSTER EXCEL PARSER")
    print("="*50)
    
    excel_path = Path("Kuccps Files/DEGREE_CLUSTER_DOCUMENT_2025_03 (1).xlsx")
    
    if not excel_path.exists():
        print(f"File not found: {excel_path}")
        return
    
    clusters = parse_cluster_excel(excel_path)
    
    # Save results
    output_path = Path("scripts/data/excel_clusters.json")
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump({
            "source": str(excel_path),
            "total_clusters": len(clusters),
            "clusters": clusters
        }, f, indent=2, ensure_ascii=False)
    
    print(f"\nSaved to: {output_path}")
    
    # Summary
    total_programs = sum(len(c["programs"]) for c in clusters)
    print(f"Total program mappings: {total_programs}")
    
    # Show sample
    print("\nSample clusters:")
    for c in clusters[:3]:
        print(f"\n  {c['cluster_id']}:")
        print(f"    S1: {c['subjects']['subject_1']} ({c['subjects']['grade_1']})")
        print(f"    S2: {c['subjects']['subject_2']} ({c['subjects']['grade_2']})")
        print(f"    Programs: {len(c['programs'])}")
    
    print("\n" + "="*50)
    print("DONE")

if __name__ == "__main__":
    main()
