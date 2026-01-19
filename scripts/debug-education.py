# Debug Education parsing
import openpyxl
import re

EXCEL_PATH = "Kuccps Files/TVET_CLUSTER_DOCUMENT_2025 - updated.xlsx"

def clean_grade(grade_str):
    text = str(grade_str).lower()
    text = re.sub(r'\(plain\)|plain', '', text)
    text = re.sub(r'\(minus\)', '-', text)
    text = re.sub(r'\(plus\)', '+', text)
    text = text.replace(' ', '').strip()
    match = re.match(r'^([A-E][+\-]?)', text, re.IGNORECASE)
    return match.group(1).upper() if match else text.upper()

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb.active

# Find Education row (Row 2 in data, likely row 6-7 in sheet)
for r in range(1, 20):
    val = ws.cell(row=r, column=2).value
    if val and 'Education' in str(val):
        print(f"Found Education at Row {r}")
        reqs = ws.cell(row=r, column=5).value
        print(f"Raw Content:\n{reqs!r}")
        
        # Test Regex
        print("\nRegex Testing:")
        clean_text = str(reqs).replace('\n', '  ').strip()
        print(f"Clean Text: {clean_text}")
        
        # Regex from previous script
        pattern = r'([A-Z][A-Za-z0-9/\s&.]*?(?:Alternative\s*[A-B])?)\s*[-–:]\s*([A-E](?:[+\-]|(?:\s*\(plain\))|(?:\s*\(minus\))|(?:\s*\(plus\)))?)'
        
        matches = re.finditer(pattern, clean_text)
        for m in matches:
            subj = m.group(1).strip()
            grd = m.group(2).strip()
            print(f"MATCH: '{subj}' -> '{grd}'")
