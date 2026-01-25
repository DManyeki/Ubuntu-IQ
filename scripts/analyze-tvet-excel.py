# Analyze TVET Excel file structure
# Usage: python scripts/analyze-tvet-excel.py

import openpyxl
import json

EXCEL_PATH = "Kuccps Files/TVET_CLUSTER_DOCUMENT_2025.xlsx"
OUTPUT_PATH = "scripts/data/tvet_excel_analysis.txt"

def analyze_excel():
    print("TVET EXCEL ANALYZER")
    print("=" * 50)
    
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    print(f"Sheets: {wb.sheetnames}")
    print(f"Total sheets: {len(wb.sheetnames)}")
    
    output = []
    output.append(f"TVET Excel Analysis\n{'='*50}\n")
    output.append(f"File: {EXCEL_PATH}\n")
    output.append(f"Total sheets: {len(wb.sheetnames)}\n\n")
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        output.append(f"=== {sheet_name} ===\n")
        output.append(f"Rows: {ws.max_row}, Columns: {ws.max_column}\n")
        
        # Get first 15 rows to understand structure
        output.append("\nFirst 15 rows:\n")
        for row_idx in range(1, min(16, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                val = str(cell.value) if cell.value else ""
                if len(val) > 50:
                    val = val[:47] + "..."
                row_data.append(val)
            output.append(f"  Row {row_idx}: {' | '.join(row_data)}\n")
        
        output.append("\n")
        print(f"Analyzed sheet: {sheet_name} ({ws.max_row} rows x {ws.max_column} cols)")
    
    # Save analysis
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        f.writelines(output)
    
    print(f"\nSaved analysis to: {OUTPUT_PATH}")
    
    # Also print to console
    print("\n" + "".join(output[:100]))

if __name__ == "__main__":
    analyze_excel()
