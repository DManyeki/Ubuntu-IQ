# Final TVET Excel Extraction
# Handles merged cells, complex requirements, and schema from screenshot
# Usage: python scripts/extract-tvet-final.py

import openpyxl
import json
import re

EXCEL_PATH = "Kuccps Files/TVET_CLUSTER_DOCUMENT_2025 - updated.xlsx"
OUTPUT_PATH = "scripts/data/tvet_final_extracted.json"

def clean_grade(grade_str):
    """Normalize grade format to simple letters (C, C-, D+)"""
    if not grade_str:
        return ""
    
    # Remove "plain", "(plain)", "(minus)", "(plus)"
    text = str(grade_str).lower()
    text = re.sub(r'\(plain\)|plain', '', text)
    text = re.sub(r'\(minus\)', '-', text)
    text = re.sub(r'\(plus\)', '+', text)
    text = text.replace(' ', '').strip()
    
    # Extract the grade part (e.g., "C+" from "C+ in two teaching subjects")
    match = re.match(r'^([A-E][+\-]?)', text, re.IGNORECASE)
    if match:
        return match.group(1).upper()
    
    return text.upper()

def parse_requirements(req_text):
    """Parse flattened requirements text into structured list"""
    if not req_text:
        return []
    
    text = str(req_text).strip()
    if text.lower() in ['none', 'n/a', '-']:
        return []
    
    requirements = []
    
    # Clean up the text to make parsing easier
    # First, strip known headers that mess up regex
    remove_headers = [
        "Science Based Courses", "Non-Science Based Courses", 
        "Visually and Hearing Impaired Applicants",
        "Minimum Subject Grade", "Subjects Requirements",
        "two teaching subjects"
    ]
    for header in remove_headers:
        text = re.sub(header, " ", text, flags=re.IGNORECASE)

    # Replace newlines with spaces to handle wrapped lines
    clean_text = text.replace('\n', '  ').strip()
    
    # Regex to find all Subject-Grade pairs
    # Improved checks: Subject must be reasonably short (not a sentence)
    pattern = r'([A-Z][A-Z0-9/&.]*(?:\s+[A-Z0-9/&.]+){0,4}(?:\s*Alternative\s*[A-B])?)\s*[-–:]\s*([A-E](?:[+\-]|(?:\s*\(plain\))|(?:\s*\(minus\))|(?:\s*\(plus\)))?)'
    
    matches = re.finditer(pattern, clean_text)
    
    for match in matches:
        subject = match.group(1).strip()
        grade_raw = match.group(2).strip()
        
        # Filter out short noise or remaining headers
        if len(subject) < 2 or "based" in subject.lower():
            continue
            
        grade = clean_grade(grade_raw)
        
        if subject and grade:
            requirements.append({
                "subject": subject,
                "grade": grade
            })
            
    return requirements

def extract_tvet_data():
    print("TVET FINAL EXTRACTION")
    print("=" * 50)
    
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active
    
    print(f"Sheet: {ws.title} ({ws.max_row} rows)")
    
    programs = []
    last_category = None
    
    # Skip header rows - assuming data starts at row 4 based on analysis
    # Column mapping based on screenshot:
    # Col 1: Number
    # Col 2: Category
    # Col 3: Level
    # Col 4: Mean Grade
    # Col 5: Requirements
    
    start_row = 1
    # Find data start (first row with a number in col 1)
    for r in range(1, 10):
        val = ws.cell(row=r, column=1).value
        if val and str(val).strip().isdigit():
            start_row = r
            break
            
    print(f"Data starts at row {start_row}")
    
    for row_idx in range(start_row, ws.max_row + 1):
        col_1_num = ws.cell(row=row_idx, column=1).value
        col_2_cat = ws.cell(row=row_idx, column=2).value
        col_3_level = ws.cell(row=row_idx, column=3).value
        col_4_grade = ws.cell(row=row_idx, column=4).value
        col_5_reqs = ws.cell(row=row_idx, column=5).value
        
        # Handle Merged Category: If col 2 is empty but col 3 has level, use last category
        category = str(col_2_cat).strip() if col_2_cat else None
        
        if category:
            last_category = category
        elif col_3_level and last_category:
            # Continuation of previous category (merged cell behavior)
            category = last_category
        else:
            # Empty row or spacer
            continue
            
        # Clean Level
        level_raw = str(col_3_level).strip() if col_3_level else ""
        if not level_raw: 
            continue
            
        level = "Unknown"
        if "Diploma" in level_raw and "Visually" in level_raw:
            level = "Diploma (Special Needs)"
        elif "Diploma" in level_raw:
            level = "Diploma"
        elif "Certificate" in level_raw:
            level = "Certificate"
        elif "Artisan" in level_raw:
            level = "Artisan"
        else:
            # Skip rows that are just headers repeated or notes
            continue
            
        # Mean Grade
        min_mean_grade = clean_grade(col_4_grade)
        
        # Requirements
        requirements = parse_requirements(col_5_reqs)
        
        # Create program entry
        program = {
            "category": category.replace('\n', ' ').strip(),
            "level": level,
            "min_mean_grade": min_mean_grade,
            "requirements": requirements
        }
        
        programs.append(program)
        print(f"Found: {program['category']} ({program['level']}) - {program['min_mean_grade']}")

    # Validation
    print(f"\nExtracted {len(programs)} programs")
    
    # Save
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump({
            "source": "TVET_CLUSTER_DOCUMENT_2025 - updated.xlsx",
            "count": len(programs),
            "programs": programs
        }, f, indent=2)
        
    print(f"Saved to {OUTPUT_PATH}")

if __name__ == "__main__":
    extract_tvet_data()
